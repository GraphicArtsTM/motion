"""Справка-анкета (объективка) в DOCX: Times New Roman 12 pt, фото 3×4, таблица родственников."""
import os
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

sys.path.insert(0, os.path.dirname(__file__))
from data import *  # noqa: E402,F403

OUT = sys.argv[1]
FONT = "Times New Roman"

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin, sec.right_margin = Cm(2.5), Cm(1.5)
sec.top_margin, sec.bottom_margin = Cm(1.5), Cm(1.5)

st = doc.styles["Normal"]
st.font.name, st.font.size = FONT, Pt(12)
st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
st.paragraph_format.space_after = Pt(0)
st.paragraph_format.line_spacing = 1.0


def run(p, text, bold=False):
    r = p.add_run(str(text))
    r.bold = bold
    r.font.name = FONT
    r._element.rPr.rFonts.set(qn("w:cs"), FONT)
    if isinstance(text, T):
        r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return r


def para(container, parts, align=None, after=0, before=0):
    p = container.add_paragraph()
    for part in parts:
        if isinstance(part, tuple):
            run(p, *part)
        else:
            run(p, part)
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    return p


def cell_text(cell, parts, bold=False, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    for part in parts if isinstance(parts, list) else [parts]:
        run(p, part, bold)
    if align:
        p.alignment = align


def borders(table, on=True):
    tbl = table._tbl
    pr = tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for e in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{e}")
        el.set(qn("w:val"), "single" if on else "nil")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "000000")
        b.append(el)
    pr.append(b)


def widths(table, cms):
    table.autofit = False
    for col, w in zip(table.columns, cms):
        col.width = Cm(w)
    for row in table.rows:
        for c, w in zip(row.cells, cms):
            c.width = Cm(w)


C = WD_ALIGN_PARAGRAPH.CENTER
para(doc, [("СПРАВКА-АНКЕТА", True)], C)
para(doc, ["(объективка)"], C, after=6)
para(doc, [(NAME_RU_FULL.upper(), True)], C, after=8)

# --- Шапка: поля слева, фото справа -----------------------------------------
top = doc.add_table(rows=1, cols=2)
borders(top, False)
widths(top, [13, 4])
left, right = top.rows[0].cells
left.text = ""
head = [
    ("1. Фамилия, имя, отчество: ", NAME_RU_FULL),
    ("2. Год и место рождения: ", f"{BIRTH_RU_TEXT}, {BIRTHPLACE_RU}, Республика Узбекистан"),
    ("3. Национальность: ", "узбек"),
    ("4. Партийность: ", "беспартийный"),
    ("5. Образование: ", "неоконченное высшее (студент 4-го курса бакалавриата)"),
    ("6. Специальность по образованию: ", "искусственный интеллект"),
]
for i, (k, v) in enumerate(head):
    p = left.paragraphs[0] if i == 0 else left.add_paragraph()
    run(p, k, True)
    run(p, v)
    p.paragraph_format.space_after = Pt(5)
right.text = ""
pp = right.paragraphs[0]
pp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
pp.add_run().add_picture(os.path.join(os.path.dirname(__file__), "photo_3x4.jpg"), width=Cm(3), height=Cm(4))

para(doc, [], after=4)

# --- Основные пункты ---------------------------------------------------------
school_line = [SCHOOL_RU, " Сурхандарьинской области, ", SCHOOL_YEARS[0], " – ", SCHOOL_YEARS[1],
               " (общее среднее образование)."]
rows = [
    ("7. Какие учебные заведения окончил", [
        "Ташкентский государственный экономический университет, направление «Искусственный интеллект», "
        "сентябрь 2023 года – по настоящее время (обучение продолжается, окончание — июнь 2027 года);\n",
        *school_line]),
    ("8. Учёная степень, учёное звание", ["нет"]),
    ("9. Какими иностранными языками владеет", [
        "узбекский — родной; русский — свободно (уровень C1); английский — средний уровень (B1)"]),
    ("10. Награды (государственные и иные)", ["нет"]),
    ("11. Воинское звание", ["нет"]),
    ("12. Является ли депутатом представительных органов", ["нет"]),
    ("13. Постоянное место жительства", [ADDRESS_RU + "; временно проживает в городе Ташкенте"]),
    ("14. Контактные данные", ["телефон: ", PHONE, f"; электронная почта: {EMAIL}; Telegram: {TELEGRAM}"]),
]
t = doc.add_table(rows=len(rows), cols=2)
borders(t)
widths(t, [6, 11])
for r, (k, v) in zip(t.rows, rows):
    cell_text(r.cells[0], k, bold=True)
    cell_text(r.cells[1], v)

# --- Трудовая деятельность ----------------------------------------------------
para(doc, [("15. ТРУДОВАЯ ДЕЯТЕЛЬНОСТЬ", True)], C, before=10, after=4)
work = [
    ([SCHOOL_YEARS[0], " – ", SCHOOL_YEARS[1]], ["учащийся: ", SCHOOL_RU, " Сурхандарьинской области"]),
    ([SCHOOL_YEARS[1], " – 08.2023"], [GAP_RU]),
    (["09.2023 – по настоящее время"], ["студент Ташкентского государственного экономического университета, "
                                         "направление «Искусственный интеллект»"]),
    ([FREELANCE_FROM, " – по настоящее время"], ["видеомонтажёр и моушн-дизайнер (фриланс, работа по договорённости "
                                                "с заказчиками) под брендом GraphicArts, город Ташкент"]),
]
wt = doc.add_table(rows=len(work) + 1, cols=2)
borders(wt)
widths(wt, [5, 12])
cell_text(wt.rows[0].cells[0], "Годы (месяц.год)", True, C)
cell_text(wt.rows[0].cells[1], "Место работы (учёбы) и должность", True, C)
for r, (d, w) in zip(wt.rows[1:], work):
    cell_text(r.cells[0], d)
    cell_text(r.cells[1], w)

# --- Родственники ------------------------------------------------------------
doc.add_page_break()
para(doc, [("СВЕДЕНИЯ", True)], C)
para(doc, [("о близких родственниках Рузибоева Мухаммада Ильхомовича", True)], C, after=6)
rt = doc.add_table(rows=len(RELATIVES) + 1, cols=5)
borders(rt)
cw = [2.6, 3.8, 3.4, 3.8, 3.4]
widths(rt, cw)
rt.alignment = WD_TABLE_ALIGNMENT.CENTER
for c, hname in zip(rt.rows[0].cells, ["Степень родства", "Фамилия, имя, отчество", "Год и место рождения",
                                       "Место работы и должность", "Место жительства"]):
    cell_text(c, hname, True, C)
for r, rel in zip(rt.rows[1:], RELATIVES):
    for c, v in zip(r.cells, rel):
        cell_text(c, v)

para(doc, ["Дата заполнения: «29» октября 2026 года"], before=24)
para(doc, ["Подпись: ______________________   Рузибоев Мухаммад Ильхомович"], before=14)

doc.save(os.path.join(OUT, "Рузибоев_Мухаммад_Объективка.docx"))
